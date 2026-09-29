"""Development fixtures; requires python-docx, openpyxl and python-pptx."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

output = Path(__file__).resolve().parent.parent / "artifacts" / "office"
output.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("LanPrint 文档打印验证", 0)
doc.add_paragraph("Word → PDF → 预览 → 打印。第一页。")
doc.add_page_break()
doc.add_paragraph("第二页：中文、空格文件名和页码验证。")
doc.save(output / "中文 文档.docx")
book = Workbook()
sheet = book.active
sheet.title = "打印测试"
sheet.append(["项目", "数量", "金额"])
sheet.append(["打印纸", 2, 35])
sheet.append(["总计", "=SUM(B2:B2)", "=SUM(C2:C2)"])
sheet.print_area = "A1:C3"
sheet.column_dimensions["A"].width = 25
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
book.save(output / "中文 表格.xlsx")
deck = Presentation()
for i in range(2):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = f"LanPrint 演示文稿 {i + 1}"
    slide.placeholders[1].text = "PowerPoint → PDF → 预览 → 打印"
deck.save(output / "中文 演示.pptx")
print(output)
